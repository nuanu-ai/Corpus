"""Tests for format-specific document extractor.

All tests use tempfile for file creation.
XLSX tests require openpyxl (skipped if not installed).
PDF tests require pdfplumber (skipped if not installed).
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest

from extraction.extractor import (
    ExtractionResult,
    extract,
    extract_csv,
    extract_image,
    extract_xlsx,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _write_tmp(suffix: str, content: str | bytes, tmp_dir: Path) -> str:
    """Write content to a temp file and return the path as str."""
    p = tmp_dir / f"test_file{suffix}"
    if isinstance(content, bytes):
        p.write_bytes(content)
    else:
        p.write_text(content, encoding="utf-8")
    return str(p)


# ---------------------------------------------------------------------------
# 1. CSV extraction with standard data
# ---------------------------------------------------------------------------


class TestCSVStandard:
    def test_basic_csv(self, tmp_dir: Path):
        path = _write_tmp(
            ".csv",
            "date,description,amount\n"
            "2025-01-15,Office supplies,-120.50\n"
            "2025-01-16,Client payment,5000.00\n",
            tmp_dir,
        )
        result = extract_csv(path)

        assert isinstance(result, ExtractionResult)
        assert result.format == "csv"
        assert result.row_count == 2
        assert len(result.records) == 2
        assert result.records[0]["date"] == "2025-01-15"
        assert result.records[0]["description"] == "Office supplies"
        assert result.records[0]["amount"] == "-120.50"
        assert result.records[1]["amount"] == "5000.00"
        assert result.dedup_hash  # non-empty SHA-256

    def test_headers_stripped(self, tmp_dir: Path):
        path = _write_tmp(
            ".csv",
            " date , description , amount \n"
            "2025-01-15,Test,100\n",
            tmp_dir,
        )
        result = extract_csv(path)
        assert list(result.records[0].keys()) == ["date", "description", "amount"]


# ---------------------------------------------------------------------------
# 2. CSV with semicolon delimiter
# ---------------------------------------------------------------------------


class TestCSVSemicolon:
    def test_semicolon_delimiter(self, tmp_dir: Path):
        path = _write_tmp(
            ".csv",
            "date;description;amount\n"
            "2025-01-15;Buroeinrichtung;-120,50\n"
            "2025-01-16;Kundenzahlung;5000,00\n",
            tmp_dir,
        )
        result = extract_csv(path)

        assert result.format == "csv"
        assert result.row_count == 2
        assert result.records[0]["date"] == "2025-01-15"
        assert result.records[0]["description"] == "Buroeinrichtung"
        assert result.records[0]["amount"] == "-120,50"

    def test_tab_delimiter(self, tmp_dir: Path):
        path = _write_tmp(
            ".csv",
            "date\tdescription\tamount\n"
            "2025-01-15\tRent\t-2000\n",
            tmp_dir,
        )
        result = extract_csv(path)

        assert result.row_count == 1
        assert result.records[0]["description"] == "Rent"


# ---------------------------------------------------------------------------
# 3. CSV with BOM
# ---------------------------------------------------------------------------


class TestCSVBom:
    def test_utf8_bom(self, tmp_dir: Path):
        bom_content = b"\xef\xbb\xbfdate,description,amount\n2025-01-15,Test,100\n"
        path = _write_tmp(".csv", bom_content, tmp_dir)
        result = extract_csv(path)

        assert result.row_count == 1
        # The key must be "date", not "\ufeffdate"
        assert "date" in result.records[0]
        assert result.records[0]["date"] == "2025-01-15"


# ---------------------------------------------------------------------------
# 4. XLSX extraction (skip if openpyxl unavailable)
# ---------------------------------------------------------------------------


class TestXLSX:
    @pytest.fixture(autouse=True)
    def _require_openpyxl(self):
        pytest.importorskip("openpyxl")

    def _create_xlsx(self, tmp_dir: Path, rows: list[list], sheet_name: str = "Sheet1") -> str:
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in rows:
            ws.append(row)
        path = str(tmp_dir / "test.xlsx")
        wb.save(path)
        wb.close()
        return path

    def test_basic_xlsx(self, tmp_dir: Path):
        path = self._create_xlsx(tmp_dir, [
            ["date", "description", "amount"],
            ["2025-01-15", "Office supplies", -120.50],
            ["2025-01-16", "Client payment", 5000.00],
        ])
        result = extract_xlsx(path)

        assert result.format == "xlsx"
        assert result.row_count == 2
        assert result.records[0]["date"] == "2025-01-15"
        assert result.records[0]["amount"] == -120.50

    def test_skips_empty_rows(self, tmp_dir: Path):
        path = self._create_xlsx(tmp_dir, [
            ["name", "value"],
            ["Alice", 100],
            [None, None],
            ["Bob", 200],
        ])
        result = extract_xlsx(path)

        assert result.row_count == 2
        assert result.records[0]["name"] == "Alice"
        assert result.records[1]["name"] == "Bob"

    def test_sheet_name_selection(self, tmp_dir: Path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Summary"
        ws1.append(["key", "value"])
        ws1.append(["total", 999])

        ws2 = wb.create_sheet("Details")
        ws2.append(["item", "cost"])
        ws2.append(["Widget", 42])

        path = str(tmp_dir / "multi_sheet.xlsx")
        wb.save(path)
        wb.close()

        result = extract_xlsx(path, sheet_name="Details")
        assert result.row_count == 1
        assert result.records[0]["item"] == "Widget"

    def test_missing_sheet_name(self, tmp_dir: Path):
        path = self._create_xlsx(tmp_dir, [["a"], [1]])
        result = extract_xlsx(path, sheet_name="NonExistent")
        assert result.row_count == 0
        assert any("not found" in w for w in result.warnings)

    def test_merged_cells(self, tmp_dir: Path):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["category", "item", "amount"])
        ws.append(["Food", "Pizza", 15])
        ws.append(["Food", "Burger", 10])
        # Merge category cells A2:A3
        ws.merge_cells("A2:A3")

        path = str(tmp_dir / "merged.xlsx")
        wb.save(path)
        wb.close()

        result = extract_xlsx(path)
        assert result.row_count == 2
        # Both rows should have the merged value "Food"
        assert result.records[0]["category"] == "Food"
        assert result.records[1]["category"] == "Food"


# ---------------------------------------------------------------------------
# 5. Empty document -> empty records
# ---------------------------------------------------------------------------


class TestEmptyDocument:
    def test_empty_csv(self, tmp_dir: Path):
        path = _write_tmp(".csv", "", tmp_dir)
        result = extract_csv(path)

        assert result.row_count == 0
        assert result.records == []
        assert any("Empty" in w for w in result.warnings)

    def test_empty_csv_whitespace_only(self, tmp_dir: Path):
        path = _write_tmp(".csv", "   \n  \n", tmp_dir)
        result = extract_csv(path)

        assert result.row_count == 0
        assert result.records == []

    def test_csv_headers_only(self, tmp_dir: Path):
        path = _write_tmp(".csv", "date,description,amount\n", tmp_dir)
        result = extract_csv(path)

        assert result.row_count == 0
        assert result.records == []


# ---------------------------------------------------------------------------
# 6. Dedup hash consistency
# ---------------------------------------------------------------------------


class TestDedupHash:
    def test_same_file_same_hash(self, tmp_dir: Path):
        content = "date,description,amount\n2025-01-15,Test,100\n"
        path = _write_tmp(".csv", content, tmp_dir)

        result1 = extract_csv(path)
        result2 = extract_csv(path)

        assert result1.dedup_hash == result2.dedup_hash
        assert len(result1.dedup_hash) == 64  # SHA-256 hex digest length

    def test_different_files_different_hash(self, tmp_dir: Path):
        path1 = _write_tmp(".csv", "a,b\n1,2\n", tmp_dir)
        p2 = tmp_dir / "other.csv"
        p2.write_text("x,y\n9,8\n", encoding="utf-8")

        result1 = extract_csv(path1)
        result2 = extract_csv(str(p2))

        assert result1.dedup_hash != result2.dedup_hash


# ---------------------------------------------------------------------------
# 7. Auto-detect format by extension
# ---------------------------------------------------------------------------


class TestAutoDetect:
    def test_csv_auto(self, tmp_dir: Path):
        path = _write_tmp(".csv", "a,b\n1,2\n", tmp_dir)
        result = extract(path)
        assert result.format == "csv"
        assert result.row_count == 1

    def test_image_auto(self, tmp_dir: Path):
        # Create a minimal file with .png extension (content doesn't matter for stub)
        path = _write_tmp(".png", b"\x89PNG\r\n\x1a\n", tmp_dir)
        result = extract(path)
        assert result.format == "image"
        assert result.row_count == 0
        assert any("OCR" in w or "Vision" in w for w in result.warnings)

    def test_unsupported_extension(self, tmp_dir: Path):
        path = _write_tmp(".xyz", "hello", tmp_dir)
        with pytest.raises(ValueError, match="Unsupported file format"):
            extract(path)

    def test_xlsx_auto(self, tmp_dir: Path):
        openpyxl = pytest.importorskip("openpyxl")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["col1", "col2"])
        ws.append([10, 20])
        path = str(tmp_dir / "auto.xlsx")
        wb.save(path)
        wb.close()

        result = extract(path)
        assert result.format == "xlsx"
        assert result.row_count == 1

    def test_jpg_auto(self, tmp_dir: Path):
        path = _write_tmp(".jpg", b"\xff\xd8\xff", tmp_dir)
        result = extract(path)
        assert result.format == "image"


# ---------------------------------------------------------------------------
# 8. Image extraction stub
# ---------------------------------------------------------------------------


class TestImageStub:
    def test_image_returns_empty_with_warning(self, tmp_dir: Path):
        path = _write_tmp(".png", b"\x89PNG\r\n\x1a\n", tmp_dir)
        result = extract_image(path)

        assert result.format == "image"
        assert result.row_count == 0
        assert result.records == []
        assert len(result.warnings) == 1
        assert "OCR" in result.warnings[0] or "Vision" in result.warnings[0]
        assert result.dedup_hash  # still computes a hash


# ---------------------------------------------------------------------------
# PDF extraction (skip if pdfplumber unavailable)
# ---------------------------------------------------------------------------


class TestPDF:
    @pytest.fixture(autouse=True)
    def _require_pdfplumber(self):
        pytest.importorskip("pdfplumber")

    def test_pdf_auto_detect(self, tmp_dir: Path):
        """Verify that extract() routes .pdf to extract_pdf.

        We test this indirectly: a valid PDF that pdfplumber can open
        should return format='pdf'. We create a minimal PDF using
        reportlab if available, otherwise skip.
        """
        reportlab = pytest.importorskip("reportlab")
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        path = str(tmp_dir / "test.pdf")
        c = canvas.Canvas(path, pagesize=letter)
        c.drawString(100, 700, "Hello World")
        c.save()

        result = extract(path)
        assert result.format == "pdf"
        assert result.dedup_hash
        assert len(result.dedup_hash) == 64
